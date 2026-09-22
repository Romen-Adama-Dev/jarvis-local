from pathlib import Path

import pytest
from docx import Document as DocxDocument
from pptx import Presentation

from packages.core.errors import ProviderUnavailableError
from packages.docgen.render_docx import render_docx
from packages.docgen.render_markdown import render_markdown
from packages.docgen.render_pdf import render_pdf_via_pandoc
from packages.docgen.render_pptx import render_pptx
from packages.docgen.schema import DocSection, GeneratedDoc
from packages.rag.orchestrator import RagSource


def _fake_doc() -> GeneratedDoc:
    grounded = DocSection(
        title="Fortalezas",
        answer="El equipo tiene experiencia previa en proyectos similares.",
        sources=[
            RagSource(
                document_id="doc-1",
                filename="acta.pdf",
                chunk_id="chunk-1",
                page=3,
                section=None,
                score=0.5,
            )
        ],
        insufficient_evidence=False,
    )
    ungrounded = DocSection(
        title="Amenazas",
        answer="",
        sources=[],
        insufficient_evidence=True,
    )
    return GeneratedDoc(
        kind="dafo",
        topic="Proyecto Fénix",
        generated_at="2026-09-09T00:00:00+00:00",
        sections=[grounded, ungrounded],
    )


def test_render_markdown_contains_titles_and_no_evidence_phrase():
    text = render_markdown(_fake_doc())
    assert "## Fortalezas" in text
    assert "## Amenazas" in text
    assert "acta.pdf" in text
    assert "Sin evidencia suficiente en la documentación indexada." in text


def test_render_docx_produces_readable_file(tmp_path: Path):
    out_path = tmp_path / "doc.docx"
    render_docx(_fake_doc(), out_path)

    assert out_path.exists()
    assert out_path.stat().st_size > 0

    document = DocxDocument(str(out_path))
    all_text = "\n".join(p.text for p in document.paragraphs)
    assert "Proyecto Fénix" in all_text
    assert "Fortalezas" in all_text
    assert "Amenazas" in all_text
    assert "Sin evidencia suficiente" in all_text


def test_render_pptx_produces_readable_file(tmp_path: Path):
    out_path = tmp_path / "doc.pptx"
    render_pptx(_fake_doc(), out_path)

    assert out_path.exists()
    assert out_path.stat().st_size > 0

    presentation = Presentation(str(out_path))
    assert len(presentation.slides) == 3  # title slide + 2 sections
    titles = [slide.shapes.title for slide in presentation.slides]
    assert all(title is not None for title in titles)
    assert [title.text for title in titles if title is not None] == [
        "Proyecto Fénix",
        "Fortalezas",
        "Amenazas",
    ]


def test_render_pdf_via_pandoc_raises_when_missing_binaries(tmp_path, monkeypatch):
    monkeypatch.setattr("shutil.which", lambda _name: None)
    markdown_path = tmp_path / "source.md"
    markdown_path.write_text("# hola", encoding="utf-8")

    with pytest.raises(ProviderUnavailableError):
        render_pdf_via_pandoc(markdown_path, tmp_path / "out.pdf")


def test_render_pdf_ignores_raw_latex(tmp_path: Path):
    import shutil

    if shutil.which("pandoc") is None or shutil.which("xelatex") is None:
        pytest.skip("pandoc/xelatex no instalados")
    secret = tmp_path / "secreto.txt"
    secret.write_text("MARCADOR-SECRETO")
    markdown = tmp_path / "doc.md"
    markdown.write_text(f"# Informe\n\n- informe.docx (\\input{{{secret}}})\n")
    out = tmp_path / "doc.pdf"
    render_pdf_via_pandoc(markdown, out)
    from pypdf import PdfReader

    text = "".join(page.extract_text() for page in PdfReader(str(out)).pages)
    assert "MARCADOR-SECRETO" not in text
