import importlib.util
from pathlib import Path

import pytest
from openpyxl import load_workbook

_SERVER = Path(__file__).resolve().parents[2] / "integrations/openclaw/skills/jarvis-pm/server.py"


@pytest.fixture
def pm_server(monkeypatch, tmp_path):
    spec = importlib.util.spec_from_file_location("jarvis_pm_server", _SERVER)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    monkeypatch.setattr(module, "JARVIS_OUTBOX_DIR", tmp_path)
    return module


class _FakeOpenProject:
    def find_project(self, name: str) -> dict:
        return {"id": 9, "name": "App de reservas"}

    def work_packages(self, project: dict, **kwargs) -> list[dict]:
        self.kwargs = kwargs
        return [
            {
                "id": 41,
                "subject": '=HYPERLINK("http://example.invalid")',
                "startDate": "2026-09-01",
                "dueDate": "2026-09-30",
                "percentageDone": 40,
                "description": {"raw": "Maquetar la agenda"},
                "_links": {
                    "type": {"title": "Tarea"},
                    "status": {"title": "En curso"},
                    "assignee": {"title": "Ana Ejemplo"},
                },
            }
        ]

    def work_package_url(self, wp: dict) -> str:
        return f"https://openproject.example/work_packages/{wp['id']}"


def test_pm_export_tasks_writes_xlsx_with_literal_subjects(pm_server, monkeypatch, tmp_path):
    fake = _FakeOpenProject()
    monkeypatch.setattr(pm_server, "_op", lambda: fake)
    reply = pm_server.pm_export_tasks("app de reservas")
    media = [line for line in reply.splitlines() if line.startswith("MEDIA:")]
    assert len(media) == 1
    path = Path(media[0].removeprefix("MEDIA:"))
    assert path.parent == tmp_path and path.suffix == ".xlsx"
    assert fake.kwargs["only_open"] is False
    ws = load_workbook(path)["Tareas"]
    assert [c.value for c in ws[1]][:4] == ["ID", "Tipo", "Asunto", "Estado"]
    assert ws["A2"].value == 41 and ws["H2"].value == 40
    # El asunto viene de OpenProject: se guarda como texto, nunca como fórmula.
    assert ws["C2"].data_type == "s"


def test_pm_export_tasks_rejects_unknown_format(pm_server, monkeypatch):
    monkeypatch.setattr(pm_server, "_op", lambda: _FakeOpenProject())
    assert pm_server.pm_export_tasks("app", format="pdf").startswith("No se pudo")
