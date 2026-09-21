import zipfile
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from odf import opendocument, table, text
from openpyxl import load_workbook
from pptx import Presentation

from packages.core.errors import ValidationFailedError
from packages.office import build_document, normalize_format, parse_blocks, safe_filename
from packages.office.spec import Formula, bold_runs, coerce_cell, formula_for_ods

_BLOCKS = [
    {"type": "heading", "text": "Resumen", "level": 1},
    {"type": "paragraph", "text": "Presupuesto **aprobado** por el comité."},
    {"type": "bullets", "items": ["Fase 1 cerrada", "Fase 2 en **curso**"]},
    {
        "type": "table",
        "name": "Costes",
        "columns": ["Partida", "Horas", "Importe"],
        "rows": [["Diseño", 10, "1200,50"], ["Desarrollo", "25", 3000], ["Pruebas", 5, 600]],
        "total": True,
    },
]


def test_bold_runs_splits_markdown_bold():
    assert bold_runs("a **b** c") == [("a ", False), ("b", True), (" c", False)]


def test_coerce_cell_numbers_formulas_and_literals():
    assert coerce_cell("12,5", allow_formulas=True) == 12.5
    assert coerce_cell("7", allow_formulas=True) == 7
    assert coerce_cell("=SUM(B2:B4)", allow_formulas=True) == Formula("=SUM(B2:B4)")
    # DDE, referencias externas y datos de fuera nunca son fórmulas.
    assert coerce_cell("=cmd|' /C calc'!A0", allow_formulas=True) == "=cmd|' /C calc'!A0"
    assert coerce_cell("=Hoja2!A1", allow_formulas=True) == "=Hoja2!A1"
    assert coerce_cell("=SUM(B2:B4)", allow_formulas=False) == "=SUM(B2:B4)"
    assert coerce_cell("2026-09-24", allow_formulas=True) == "2026-09-24"


def test_formula_for_ods():
    assert formula_for_ods(Formula("=SUM(B2:B5)")) == "of:=SUM([.B2:.B5])"
    assert formula_for_ods(Formula("=ROUND(C2,2)")) == "of:=ROUND([.C2];2)"


def test_parse_blocks_rejects_bad_input():
    with pytest.raises(ValidationFailedError):
        parse_blocks([])
    with pytest.raises(ValidationFailedError):
        parse_blocks([{"type": "imagen"}])
    with pytest.raises(ValidationFailedError):
        parse_blocks([{"type": "table", "columns": [], "rows": []}])


def test_normalize_format_and_filename():
    assert normalize_format("Excel") == "xlsx"
    assert normalize_format(".ODT") == "odt"
    with pytest.raises(ValidationFailedError):
        normalize_format("pdf")
    assert safe_filename("../Acta Reunión 1/2") == "acta-reunion-1-2"


def test_build_document_strips_extension_from_filename(tmp_path: Path):
    blocks = parse_blocks([{"type": "paragraph", "text": "x"}])
    out = build_document("docx", "Acta", blocks, tmp_path, filename="acta.docx")
    assert out.name.startswith("acta-2") and out.suffix == ".docx"


def test_xlsx_has_formatted_header_numbers_and_total(tmp_path: Path):
    out = build_document("xlsx", "Costes del proyecto", parse_blocks(_BLOCKS), tmp_path)
    assert out.parent == tmp_path and out.suffix == ".xlsx"
    wb = load_workbook(out)
    assert wb.sheetnames == ["Costes", "Notas"]
    ws = wb["Costes"]
    assert ws["A1"].value == "Partida" and ws["A1"].font.bold
    assert ws["C2"].value == 1200.5 and ws["B3"].value == 25
    assert ws["A5"].value == "Total" and ws["C5"].value == "=SUM(C2:C4)"
    assert ws.freeze_panes == "A2"
    assert "Presupuesto aprobado por el comité." in [c.value for c in wb["Notas"]["A"]]


def test_xlsx_keeps_untrusted_equals_as_text(tmp_path: Path):
    blocks = parse_blocks(
        [{"type": "table", "columns": ["Asunto"], "rows": [["=HYPERLINK(1)"]]}],
        allow_formulas=False,
    )
    ws = load_workbook(build_document("xlsx", "T", blocks, tmp_path)).active
    assert ws is not None
    assert ws["A2"].data_type == "s"


def test_docx_has_bold_runs_and_table(tmp_path: Path):
    out = build_document("docx", "Acta", parse_blocks(_BLOCKS), tmp_path)
    doc = DocxDocument(str(out))
    bold = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
    assert "aprobado" in bold and "curso" in bold
    grid = doc.tables[0]
    assert [c.text for c in grid.rows[0].cells] == ["Partida", "Horas", "Importe"]
    assert grid.rows[-1].cells[0].text == "Total"
    assert grid.rows[-1].cells[1].text == "40"
    assert grid.rows[2].cells[2].text == "3.000"


def test_pptx_has_cover_content_and_table_slides(tmp_path: Path):
    out = build_document("pptx", "Estado", parse_blocks(_BLOCKS), tmp_path)
    prs = Presentation(str(out))
    titles = [s.shapes.title.text for s in prs.slides if s.shapes.title is not None]
    assert titles == ["Estado", "Resumen", "Costes"]
    assert any(shape.has_table for shape in prs.slides[2].shapes)


@pytest.mark.parametrize("fmt", ["ods", "odt", "odp"])
def test_opendocument_files_are_valid(tmp_path: Path, fmt: str):
    out = build_document(fmt, "Plan", parse_blocks(_BLOCKS), tmp_path)
    with zipfile.ZipFile(out) as zf:
        assert (
            zf.read("mimetype")
            .decode()
            .endswith({"ods": "spreadsheet", "odt": "text", "odp": "presentation"}[fmt])
        )
    doc = opendocument.load(str(out))
    body = "".join(str(p) for p in doc.getElementsByType(text.P))
    assert "Presupuesto" in body
    if fmt == "ods":
        cells = doc.getElementsByType(table.TableCell)
        formulas = [c.getAttribute("formula") for c in cells if c.getAttribute("formula")]
        assert formulas == ["of:=SUM([.B2:.B4])", "of:=SUM([.C2:.C4])"]
