"""Renderizadores Office Open XML: Excel (openpyxl), Word (python-docx) y PowerPoint
(python-pptx)."""

from pathlib import Path

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from pptx import Presentation
from pptx.util import Inches, Pt

from packages.office.spec import (
    Block,
    Cell,
    Formula,
    bold_runs,
    column_letter,
    column_sum,
    numeric_columns,
    plain,
)

_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_TOTAL_FONT = Font(bold=True)
_MAX_WIDTH = 60


def display(value: Cell | Formula) -> str:
    if value is None:
        return ""
    if isinstance(value, Formula):
        return value.text
    if isinstance(value, float):
        return f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    if isinstance(value, int):
        return f"{value:,}".replace(",", ".")
    return str(value)


def _tables(blocks: list[Block]) -> list[Block]:
    return [b for b in blocks if b.kind == "table"]


def _notes(blocks: list[Block]) -> list[str]:
    """Texto que no es tabla, para la hoja «Notas» de un libro de cálculo."""
    lines: list[str] = []
    for block in blocks:
        if block.kind in ("heading", "paragraph"):
            lines.append(plain(block.text))
        elif block.kind == "bullets":
            lines.extend(f"• {plain(item)}" for item in block.items)
    return lines


def sheet_names(blocks: list[Block]) -> list[str]:
    """Nombres únicos de hoja (máx. 31 caracteres, sin los prohibidos por Excel)."""
    names: list[str] = []
    for index, table in enumerate(_tables(blocks), start=1):
        base = "".join(c for c in (table.name or f"Hoja{index}") if c not in "[]:*?/\\")
        base = base.strip()[:31] or f"Hoja{index}"
        name, n = base, 2
        while name.lower() in (x.lower() for x in names):
            suffix = f" ({n})"
            name, n = base[: 31 - len(suffix)] + suffix, n + 1
        names.append(name)
    return names


def total_row(block: Block, numeric: set[int]) -> list[Cell | Formula]:
    row: list[Cell | Formula] = ["Total"]
    row += [column_sum(block, c) if c in numeric else None for c in range(1, len(block.columns))]
    return row


# --- Excel --------------------------------------------------------------------------


def _write_table(ws: Worksheet, table: Block) -> None:
    ws.append(table.columns)
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r, row in enumerate(table.rows, start=2):
        for col, value in enumerate(row, start=1):
            written = value.text if isinstance(value, Formula) else value
            cell = ws.cell(row=r, column=col, value=written)
            # Un texto que empieza por «=» es literal, nunca una fórmula.
            if isinstance(value, str) and value.startswith("="):
                cell.data_type = "s"
            if isinstance(value, float):
                cell.number_format = "#,##0.00"
            elif isinstance(value, str) and len(value) > _MAX_WIDTH:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    if table.total and table.rows:
        last = ws.max_row
        total_row = last + 1
        ws.cell(row=total_row, column=1, value="Total").font = _TOTAL_FONT
        for col in numeric_columns(table):
            if col == 0:
                continue
            letter = column_letter(col)
            cell = ws.cell(row=total_row, column=col + 1, value=f"=SUM({letter}2:{letter}{last})")
            cell.font = _TOTAL_FONT
            cell.number_format = "#,##0.00"
    ws.freeze_panes = "A2"
    if table.rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(table.columns))}{len(table.rows) + 1}"
    for col in range(len(table.columns)):
        values = [table.columns[col]] + [display(r[col]) for r in table.rows]
        width = min(_MAX_WIDTH, max(len(v) for v in values) + 2)
        ws.column_dimensions[get_column_letter(col + 1)].width = max(8, width)


def render_xlsx(title: str, blocks: list[Block], out_path: Path) -> None:
    workbook = Workbook()
    workbook.properties.title = title
    default = workbook.active
    tables = _tables(blocks)
    for name, table in zip(sheet_names(blocks), tables, strict=True):
        _write_table(workbook.create_sheet(name), table)
    notes = _notes(blocks)
    if notes or not tables:
        ws = workbook.create_sheet("Notas", 0 if not tables else None)
        ws.append([title])
        ws["A1"].font = Font(bold=True, size=14)
        for line in notes:
            ws.append([line])
        ws.column_dimensions["A"].width = 100
    if default is not None:
        workbook.remove(default)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out_path)


# --- Word ---------------------------------------------------------------------------


def _add_runs(paragraph, text: str) -> None:
    for chunk, is_bold in bold_runs(text):
        run = paragraph.add_run(chunk)
        run.bold = is_bold or None


def render_docx(title: str, blocks: list[Block], out_path: Path) -> None:
    document = DocxDocument()
    document.core_properties.title = title
    document.add_heading(title, level=0)
    for block in blocks:
        if block.kind == "heading":
            document.add_heading(plain(block.text), level=block.level)
        elif block.kind == "paragraph":
            _add_runs(document.add_paragraph(), block.text)
        elif block.kind == "bullets":
            for item in block.items:
                _add_runs(document.add_paragraph(style="List Bullet"), item)
        elif block.kind == "table":
            _docx_table(document, block)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out_path))


def _docx_table(document, block: Block) -> None:
    if block.name:
        caption = document.add_paragraph()
        caption.add_run(block.name).bold = True
    numeric = set(numeric_columns(block))
    table = document.add_table(rows=1, cols=len(block.columns))
    table.style = "Light Grid Accent 1"
    for cell, header in zip(table.rows[0].cells, block.columns, strict=True):
        cell.text = ""
        cell.paragraphs[0].add_run(header).bold = True
    rows: list[list[Cell | Formula]] = [list(r) for r in block.rows]
    if block.total and rows:
        rows.append(total_row(block, numeric))
    for index, values in enumerate(rows):
        cells = table.add_row().cells
        is_total = block.total and index == len(rows) - 1
        for col, (cell, value) in enumerate(zip(cells, values, strict=True)):
            paragraph = cell.paragraphs[0]
            if isinstance(value, str) and not is_total:
                _add_runs(paragraph, value)
            else:
                paragraph.add_run(display(value)).bold = is_total or None
            if col in numeric:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    document.add_paragraph()


# --- PowerPoint ---------------------------------------------------------------------

_TITLE_LAYOUT, _CONTENT_LAYOUT, _TITLE_ONLY_LAYOUT = 0, 1, 5
_MAX_TABLE_ROWS_PER_SLIDE = 12


def slides(blocks: list[Block]) -> list[tuple[str, list[Block]]]:
    """Agrupa los bloques en diapositivas: un título de nivel 1 abre una nueva y cada
    tabla va en la suya."""
    slides: list[tuple[str, list[Block]]] = []
    current_title, current = "", []
    # Un título ya usado por una tabla no vuelve a abrir una diapositiva vacía.
    title_used = False
    for block in blocks:
        if block.kind == "heading" and block.level == 1:
            if current or (current_title and not title_used):
                slides.append((current_title, current))
            current_title, current, title_used = plain(block.text), [], False
        elif block.kind == "table":
            if current:
                slides.append((current_title, current))
                current = []
            slides.append((block.name or current_title, [block]))
            title_used = True
        else:
            current.append(block)
    if current or (current_title and not title_used):
        slides.append((current_title, current))
    return slides


def render_pptx(title: str, blocks: list[Block], out_path: Path) -> None:
    presentation = Presentation()
    cover = presentation.slides.add_slide(presentation.slide_layouts[_TITLE_LAYOUT])
    if cover.shapes.title is not None:
        cover.shapes.title.text = title
    for slide_title, content in slides(blocks):
        if content and content[0].kind == "table":
            _pptx_table_slides(presentation, slide_title, content[0])
            continue
        slide = presentation.slides.add_slide(presentation.slide_layouts[_CONTENT_LAYOUT])
        if slide.shapes.title is not None:
            slide.shapes.title.text = slide_title
        body = slide.placeholders[1].text_frame  # type: ignore[attr-defined]
        first = True
        for block in content:
            lines = block.items if block.kind == "bullets" else [block.text]
            for line in lines:
                paragraph = body.paragraphs[0] if first else body.add_paragraph()
                first = False
                paragraph.level = 1 if block.kind == "heading" and block.level > 1 else 0
                for chunk, is_bold in bold_runs(line):
                    run = paragraph.add_run()
                    run.text = chunk
                    run.font.bold = is_bold or block.kind == "heading" or None
    out_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(out_path))


def _pptx_table_slides(presentation, slide_title: str, block: Block) -> None:
    rows: list[list[Cell | Formula]] = [list(r) for r in block.rows]
    if block.total and rows:
        numeric = set(numeric_columns(block))
        rows.append(total_row(block, numeric))
    chunks = [
        rows[i : i + _MAX_TABLE_ROWS_PER_SLIDE]
        for i in range(0, len(rows), _MAX_TABLE_ROWS_PER_SLIDE)
    ] or [[]]
    for page, chunk in enumerate(chunks, start=1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[_TITLE_ONLY_LAYOUT])
        if slide.shapes.title is not None:
            suffix = f" ({page}/{len(chunks)})" if len(chunks) > 1 else ""
            slide.shapes.title.text = f"{slide_title}{suffix}"
        shape = slide.shapes.add_table(
            len(chunk) + 1, len(block.columns), Inches(0.4), Inches(1.5), Inches(9.2), Inches(0.4)
        )
        table = shape.table
        for col, header in enumerate(block.columns):
            table.cell(0, col).text = header
        for r, values in enumerate(chunk, start=1):
            for col, value in enumerate(values):
                cell = table.cell(r, col)
                cell.text = plain(display(value))
                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(12)
