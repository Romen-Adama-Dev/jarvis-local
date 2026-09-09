from pathlib import Path

from docx import Document as DocxDocument

from packages.docgen.schema import DocSection, GeneratedDoc

_NO_EVIDENCE = "Sin evidencia suficiente en la documentación indexada."


def _sources_text(section: DocSection) -> str:
    if section.insufficient_evidence or not section.sources:
        return f"Fuentes: {_NO_EVIDENCE}"
    parts = []
    for source in section.sources:
        location = f"página {source.page}" if source.page else (source.section or "sin sección")
        parts.append(f"{source.filename} ({location})")
    return "Fuentes: " + "; ".join(parts)


def render_docx(doc: GeneratedDoc, out_path: Path) -> None:
    document = DocxDocument()
    document.add_heading(doc.topic, level=0)
    for section in doc.sections:
        document.add_heading(section.title, level=1)
        body = section.answer if not section.insufficient_evidence else _NO_EVIDENCE
        document.add_paragraph(body)
        sources_paragraph = document.add_paragraph()
        sources_run = sources_paragraph.add_run(_sources_text(section))
        sources_run.italic = True
    out_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out_path))
