import re
from pathlib import Path
from typing import cast

from pptx import Presentation
from pptx.shapes.placeholder import SlidePlaceholder
from pptx.util import Inches

from packages.docgen.schema import DocSection, GeneratedDoc

_NO_EVIDENCE = "Sin evidencia suficiente en la documentación indexada."

_TITLE_LAYOUT = 0
_CONTENT_LAYOUT = 1

_FOOTER_LEFT = Inches(0.5)
_FOOTER_TOP = Inches(6.7)
_FOOTER_WIDTH = Inches(9.0)
_FOOTER_HEIGHT = Inches(0.6)


def _bullets(text: str) -> list[str]:
    sentences = [s.strip() for s in re.split(r"(?<=[.!?])\s+", text) if s.strip()]
    return sentences or [text.strip()]


def _footer_sources(section: DocSection) -> str:
    if section.insufficient_evidence or not section.sources:
        return _NO_EVIDENCE
    filenames = sorted({source.filename for source in section.sources})
    return "Fuentes: " + ", ".join(filenames)


def render_pptx(doc: GeneratedDoc, out_path: Path) -> None:
    presentation = Presentation()

    title_slide = presentation.slides.add_slide(presentation.slide_layouts[_TITLE_LAYOUT])
    title_shape = title_slide.shapes.title
    if title_shape is not None:
        title_shape.text = doc.topic
    if len(title_slide.placeholders) > 1:
        subtitle = cast(SlidePlaceholder, title_slide.placeholders[1])
        subtitle.text = f"Documento generado: {doc.kind}"

    for section in doc.sections:
        slide = presentation.slides.add_slide(presentation.slide_layouts[_CONTENT_LAYOUT])
        title_shape = slide.shapes.title
        if title_shape is not None:
            title_shape.text = section.title

        body_placeholder = cast(SlidePlaceholder, slide.placeholders[1])
        text_frame = body_placeholder.text_frame
        text_frame.clear()

        bullets = [_NO_EVIDENCE] if section.insufficient_evidence else _bullets(section.answer)
        text_frame.text = bullets[0]
        for bullet in bullets[1:]:
            paragraph = text_frame.add_paragraph()
            paragraph.text = bullet

        footer = slide.shapes.add_textbox(
            _FOOTER_LEFT, _FOOTER_TOP, _FOOTER_WIDTH, _FOOTER_HEIGHT
        )
        footer.text_frame.text = _footer_sources(section)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(out_path))
