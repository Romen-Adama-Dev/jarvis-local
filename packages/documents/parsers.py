import csv
import io
import re
from dataclasses import dataclass

import openpyxl
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from pypdf import PdfReader

from packages.core.errors import ValidationFailedError

_HEADING_MAX_CHARS = 90
_HEADING_NUMBERING_RE = re.compile(
    r"^(chapter|part|section)?\s*\d+(\.\d+){0,3}\.?\s+\S", re.IGNORECASE
)


def _looks_like_heading(line: str) -> bool:
    """Heurística sin metadatos de fuente (pypdf no expone tamaño de letra):
    una línea corta, sin puntuación final de frase, que además esté numerada
    ("1.2 Purpose...") o predominantemente en mayúsculas ("PART 1"), se trata
    como encabezado de sección. Imperfecta por diseño para libros escaneados o
    con maquetación compleja (ver ítem "docling-ingest" del roadmap)."""
    stripped = line.strip()
    if not stripped or len(stripped) > _HEADING_MAX_CHARS:
        return False
    if stripped.endswith((".", ",", ";", ":")) and not _HEADING_NUMBERING_RE.match(stripped):
        return False
    if _HEADING_NUMBERING_RE.match(stripped):
        return True
    letters = [c for c in stripped if c.isalpha()]
    return len(letters) >= 4 and sum(1 for c in letters if c.isupper()) / len(letters) > 0.9


@dataclass(frozen=True, slots=True)
class ParsedBlock:
    text: str
    page: int | None
    section: str | None
    order: int


@dataclass(frozen=True, slots=True)
class ParsedDocument:
    blocks: list[ParsedBlock]
    page_count: int | None


def parse_document(content: bytes, content_type: str, filename: str) -> ParsedDocument:
    if filename.lower().endswith(".md"):
        return _parse_markdown(content, filename)
    parser = _PARSERS_BY_MIME.get(content_type)
    if parser is None:
        raise ValidationFailedError(f"No hay parser disponible para el tipo {content_type}")
    return parser(content, filename)


def _parse_txt(content: bytes, filename: str) -> ParsedDocument:
    text = content.decode("utf-8", errors="replace")
    blocks = [
        ParsedBlock(text=p.strip(), page=None, section=None, order=i)
        for i, p in enumerate(text.split("\n\n"))
        if p.strip()
    ]
    return ParsedDocument(blocks=blocks, page_count=None)


def _parse_markdown(content: bytes, filename: str) -> ParsedDocument:
    text = content.decode("utf-8", errors="replace")
    blocks: list[ParsedBlock] = []
    current_section: str | None = None
    order = 0
    for paragraph in text.split("\n\n"):
        cleaned = paragraph.strip()
        if not cleaned:
            continue
        first_line = cleaned.lstrip("\n").split("\n", 1)[0]
        if first_line.lstrip().startswith("#"):
            current_section = first_line.lstrip("#").strip()
            remainder = cleaned[len(first_line) :].strip()
            if not remainder:
                continue
            cleaned = remainder
        blocks.append(ParsedBlock(text=cleaned, page=None, section=current_section, order=order))
        order += 1
    return ParsedDocument(blocks=blocks, page_count=None)


def _parse_html(content: bytes, filename: str) -> ParsedDocument:
    soup = BeautifulSoup(content, "lxml")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()

    blocks: list[ParsedBlock] = []
    current_section: str | None = None
    order = 0
    for element in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "td"]):
        text = element.get_text(strip=True)
        if not text:
            continue
        if element.name.startswith("h"):
            current_section = text
            continue
        blocks.append(ParsedBlock(text=text, page=None, section=current_section, order=order))
        order += 1

    if not blocks:
        fallback_text = soup.get_text(separator="\n", strip=True)
        if fallback_text:
            blocks.append(ParsedBlock(text=fallback_text, page=None, section=None, order=0))

    return ParsedDocument(blocks=blocks, page_count=None)


def _parse_csv(content: bytes, filename: str) -> ParsedDocument:
    text = content.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return ParsedDocument(blocks=[], page_count=None)

    header = rows[0]
    blocks: list[ParsedBlock] = []
    for i, row in enumerate(rows[1:], start=1):
        pairs = [
            f"{h.strip()}: {v.strip()}" for h, v in zip(header, row, strict=False) if v.strip()
        ]
        if not pairs:
            continue
        blocks.append(ParsedBlock(text=" | ".join(pairs), page=None, section=f"fila {i}", order=i))
    return ParsedDocument(blocks=blocks, page_count=None)


def _parse_xlsx(content: bytes, filename: str) -> ParsedDocument:
    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    blocks: list[ParsedBlock] = []
    order = 0
    for sheet in workbook.worksheets:
        header: list[str] | None = None
        for row in sheet.iter_rows(values_only=True):
            values = ["" if v is None else str(v) for v in row]
            if not any(v.strip() for v in values):
                continue
            if header is None:
                header = values
                continue
            pairs = [
                f"{h.strip()}: {v.strip()}"
                for h, v in zip(header, values, strict=False)
                if v.strip()
            ]
            if not pairs:
                continue
            blocks.append(
                ParsedBlock(text=" | ".join(pairs), page=None, section=sheet.title, order=order)
            )
            order += 1
    return ParsedDocument(blocks=blocks, page_count=None)


def _parse_docx(content: bytes, filename: str) -> ParsedDocument:
    document = DocxDocument(io.BytesIO(content))
    blocks: list[ParsedBlock] = []
    current_section: str | None = None
    order = 0
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = paragraph.style.name if paragraph.style is not None else None
        if style_name is not None and style_name.lower().startswith("heading"):
            current_section = text
            continue
        blocks.append(ParsedBlock(text=text, page=None, section=current_section, order=order))
        order += 1

    for table_index, table in enumerate(document.tables):
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if not cells:
                continue
            blocks.append(
                ParsedBlock(
                    text=" | ".join(cells),
                    page=None,
                    section=f"tabla {table_index + 1}",
                    order=order,
                )
            )
            order += 1

    return ParsedDocument(blocks=blocks, page_count=None)


def _parse_pdf(content: bytes, filename: str) -> ParsedDocument:
    try:
        reader = PdfReader(io.BytesIO(content))
    except Exception as exc:
        raise ValidationFailedError(
            f"No se pudo leer el PDF (posiblemente corrupto): {exc}"
        ) from exc

    blocks: list[ParsedBlock] = []
    order = 0
    current_section: str | None = None
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        for paragraph in text.split("\n\n"):
            cleaned = paragraph.strip()
            if not cleaned:
                continue
            if _looks_like_heading(cleaned):
                current_section = cleaned
                continue
            blocks.append(
                ParsedBlock(text=cleaned, page=page_number, section=current_section, order=order)
            )
            order += 1

    return ParsedDocument(blocks=blocks, page_count=len(reader.pages))


_PARSERS_BY_MIME = {
    "text/plain": _parse_txt,
    "text/markdown": _parse_txt,
    "text/html": _parse_html,
    "text/csv": _parse_csv,
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": _parse_xlsx,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": _parse_docx,
    "application/pdf": _parse_pdf,
}
